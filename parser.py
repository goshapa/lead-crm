# -*- coding: utf-8 -*-
"""
Парсер Яндекс.Карт: ищет компании в Ташкенте по заданным категориям
и сохраняет в Excel только те, у которых НЕТ сайта (поле "Сайт" в карточке
организации отсутствует). Для найденных компаний сохраняются:
название, телефон(ы), адрес, категория (рубрика), рейтинг, кол-во отзывов,
ссылка на карточку в Яндекс.Картах.

Как запускать:
    pip install -r requirements.txt
    python parser.py

По умолчанию используется системный Microsoft Edge (есть в Windows 11
из коробки), поэтому `playwright install` не обязателен. Если Edge не
найден — раскомментируйте/поставьте BROWSER_CHANNEL = None в CONFIG и
один раз выполните `playwright install chromium`.

Скрипт открывает окно браузера (это необходимо: Яндекс.Карты блокируют
headless-запросы). Не сворачивайте и не закрывайте это окно во время
работы. Если Яндекс покажет капчу — решите её вручную прямо в открытом
окне и нажмите Enter в консоли, когда скрипт попросит.

Настройки — в блоке CONFIG ниже.
"""

import random
import re
import time
import urllib.parse
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError


# ==================== CONFIG ====================

# Категории компаний для поиска (можно менять/дополнять как угодно)
CATEGORIES = [
    "частная школа",
]

# Города для поиска: (id, slug, отображаемое имя), как в URL Яндекс.Карт
# вида https://yandex.ru/maps/<id>/<slug>/search/<категория>/
# Чтобы найти id/slug другого города — вбей его название в поиск на
# yandex.ru/maps, открой любую рубрику ("Отели" и т.п.) в левой панели
# и посмотри на адрес: /maps/<id>/<slug>/search/...
CITIES = [
    {"id": "10335", "slug": "tashkent", "name": "Ташкент"},
]

# Сколько максимум карточек забирать по каждой категории в каждом городе
MAX_RESULTS_PER_CATEGORY = 60

# Файл результата
OUTPUT_FILE = f"yandex_maps_tashkent_chastnaya_shkola_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

# False — показывать окно браузера (рекомендуется: так проще заметить
# и вручную пройти капчу, если Яндекс её покажет).
# ВАЖНО: Яндекс.Карты блокируют запросы из headless-браузера (отдают
# пустую страницу с текстом "limited"), поэтому headless=True работать не будет.
HEADLESS = False

# Какой браузер запускать. По умолчанию используется системный Microsoft Edge
# (он есть в Windows 11 из коробки) — со встроенным Chromium от Playwright
# Яндекс иногда тоже блокирует автоматизированный запуск даже в обычном
# (не headless) режиме. Варианты: "msedge", "chrome" (если установлен) или
# None — тогда будет использован Chromium, который ставится через
# `playwright install chromium`.
BROWSER_CHANNEL = "msedge"

# Пауза между открытием карточек организаций (сек) — чтобы не выглядеть ботом
DELAY_MIN, DELAY_MAX = 1.0, 2.2

# =================================================


def build_search_url(category: str, city: dict) -> str:
    encoded = urllib.parse.quote(category)
    slug = urllib.parse.quote(city["slug"], safe="")
    return f"https://yandex.ru/maps/{city['id']}/{slug}/search/{encoded}/"


def is_captcha(page) -> bool:
    try:
        text = page.locator("body").inner_text(timeout=1000).lower()
    except Exception:
        return False
    markers = ["подтвердите, что вы не робот", "smartcaptcha", "showcaptcha", "не робот"]
    return any(m in text for m in markers) or "showcaptcha" in page.url


def wait_out_captcha(page, label: str):
    if not is_captcha(page):
        return
    print("\n⚠ Обнаружена капча Яндекса.")
    print(f"   Открой окно браузера и реши капчу вручную ({label}).")
    input("   После того как капча пройдена и страница загрузилась — нажми Enter здесь... ")


def collect_org_links(page, category: str, city: dict, max_results: int) -> list[str]:
    url = build_search_url(category, city)
    print(f"\n== {city['name']} — категория «{category}» ==")
    print(f"Открываю: {url}")
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=30000)
    except Exception as e:
        print(f"  Не удалось открыть страницу поиска ({e}), пропускаю категорию.")
        return []
    time.sleep(1.5)
    wait_out_captcha(page, f"поиск: {category}")

    try:
        page.wait_for_selector("li.search-snippet-view, .search-list-view__list", timeout=10000)
    except PlaywrightTimeoutError:
        print("  Список результатов не загрузился (возможно, по запросу ничего не найдено).")
        return []

    links: set[str] = set()
    stable_rounds = 0
    prev_count = 0

    while len(links) < max_results and stable_rounds < 4:
        hrefs = page.eval_on_selector_all(
            'li.search-snippet-view a.link-overlay[href*="/maps/org/"]',
            "els => els.map(e => e.href)",
        )
        for h in hrefs:
            m = re.match(r"^(https?://[^/]+/maps/org/[^/]+/\d+)/?", h)
            if m:
                links.add(m.group(1) + "/")

        if len(links) == prev_count:
            stable_rounds += 1
        else:
            stable_rounds = 0
        prev_count = len(links)

        page.evaluate(
            """() => {
                const list = document.querySelector('.search-list-view__list');
                if (!list) return;
                const container = list.closest('.scroll__container') || list.parentElement;
                container.scrollTop = container.scrollHeight;
            }"""
        )
        time.sleep(0.9)

    result = list(links)[:max_results]
    print(f"  Найдено ссылок на организации: {len(result)}")
    return result


def parse_org_page(page, url: str, category: str, city: dict) -> dict | None:
    """Возвращает данные компании, если у неё НЕТ сайта, иначе None."""
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=30000)
    except PlaywrightTimeoutError:
        print(f"  Таймаут при открытии {url}, пропускаю.")
        return None
    except Exception as e:
        # временные сбои сети/браузера — пропускаем карточку, не роняем весь скрипт
        print(f"  Ошибка при открытии {url} ({e}), пропускаю.")
        time.sleep(3)
        return None

    time.sleep(1.4)
    wait_out_captcha(page, url)

    try:
        # h1[itemprop="name"] встречается в обоих вариантах разметки карточки
        # организации, которые отдаёт Яндекс (полная "orgpage" при прямом
        # переходе по ссылке и "card" при переходе внутри списка).
        page.wait_for_selector('h1[itemprop="name"]', timeout=8000)
    except PlaywrightTimeoutError:
        print(f"  Карточка не загрузилась: {url}")
        return None
    except Exception as e:
        print(f"  Ошибка при загрузке карточки {url} ({e}), пропускаю.")
        return None

    try:
        data = page.evaluate(
            """() => {
                const name = document.querySelector('h1[itemprop="name"]')?.textContent?.trim() || '';
                const website = document.querySelector('a.business-urls-view__link')?.href || null;
                const addrEl = document.querySelector('.orgpage-header-view__address')
                    || document.querySelector('.business-contacts-view__address');
                let address = '';
                if (addrEl) {
                    const title = addrEl.getAttribute('title') || addrEl.getAttribute('aria-label');
                    if (title) {
                        address = title
                            .replace(/\\s*·\\s*/g, ', ')
                            .replace(/\\s*на карте.*$/i, '');
                    } else {
                        address = addrEl.innerText.replace(/\\n+/g, ', ').trim();
                    }
                }
                const phones = [...document.querySelectorAll('span[itemprop="telephone"]')]
                    .map(s => s.textContent.trim());
                const rating = document.querySelector('.business-rating-badge-view__rating-text')?.textContent?.trim() || '';
                const categories = [...document.querySelectorAll('.orgpage-categories-info-view__link')]
                    .map(e => e.textContent.trim());
                const reviewsLink = document.querySelector('a[href*="/reviews/"]');
                const reviewsText = reviewsLink ? reviewsLink.closest('div,span,li')?.textContent || '' : '';
                return { name, website, address, phones, rating, categories, reviewsText };
            }"""
        )
    except Exception as e:
        print(f"  Ошибка при разборе карточки {url} ({e}), пропускаю.")
        return None

    if data.get("website"):
        return None  # у компании есть сайт — пропускаем по условию задачи

    reviews_match = re.search(r"(\d+)", data.get("reviewsText", "") or "")
    org_id_match = re.search(r"/org/[^/]+/(\d+)/", url)

    return {
        "city": city["name"],
        "category_query": category,
        "name": data.get("name") or "",
        "yandex_category": ", ".join(data.get("categories") or []),
        "phones": "; ".join(data.get("phones") or []),
        "address": data.get("address") or "",
        "rating": data.get("rating") or "",
        "reviews": reviews_match.group(1) if reviews_match else "",
        "org_id": org_id_match.group(1) if org_id_match else "",
        "url": url,
    }


def save_to_excel(rows: list[dict], path: str, verbose: bool = True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Без сайта"

    headers = [
        "Город",
        "Категория поиска",
        "Название",
        "Рубрика (Яндекс)",
        "Телефон(ы)",
        "Адрес",
        "Рейтинг",
        "Отзывов",
        "Ссылка на Яндекс.Карты",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        ws.append(
            [
                r["city"],
                r["category_query"],
                r["name"],
                r["yandex_category"],
                r["phones"],
                r["address"],
                r["rating"],
                r["reviews"],
                r["url"],
            ]
        )

    widths = [14, 18, 30, 26, 20, 40, 10, 10, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    if verbose:
        print(f"\nСохранено в файл: {path}")


def main():
    all_rows: list[dict] = []
    seen_org_ids: set[str] = set()

    with sync_playwright() as p:
        launch_kwargs = {"headless": HEADLESS}
        if BROWSER_CHANNEL:
            launch_kwargs["channel"] = BROWSER_CHANNEL
        try:
            browser = p.chromium.launch(**launch_kwargs)
        except Exception as e:
            if BROWSER_CHANNEL:
                print(f"Не удалось запустить браузер '{BROWSER_CHANNEL}' ({e}).")
                print("Пробую встроенный Chromium (нужно: playwright install chromium)...")
                browser = p.chromium.launch(headless=HEADLESS)
            else:
                raise
        context = browser.new_context(
            locale="ru-RU",
            viewport={"width": 1400, "height": 900},
        )
        page = context.new_page()

        for city in CITIES:
            for category in CATEGORIES:
                try:
                    links = collect_org_links(page, category, city, MAX_RESULTS_PER_CATEGORY)
                except Exception as e:
                    print(f"Не удалось собрать список для «{city['name']}» / «{category}» ({e}), пропускаю.")
                    continue

                for i, link in enumerate(links, start=1):
                    print(f"  [{i}/{len(links)}] {link}")
                    try:
                        row = parse_org_page(page, link, category, city)
                    except Exception as e:
                        # любая непредвиденная ошибка на одной карточке не должна
                        # обрывать весь прогон — пропускаем и едем дальше
                        print(f"  Непредвиденная ошибка на {link} ({e}), пропускаю.")
                        row = None
                    time.sleep(random.uniform(DELAY_MIN, DELAY_MAX))

                    if row is None:
                        continue

                    if row["org_id"] and row["org_id"] in seen_org_ids:
                        continue
                    if row["org_id"]:
                        seen_org_ids.add(row["org_id"])

                    all_rows.append(row)
                    print(f"      -> без сайта, сохранено: {row['name']}")
                    save_to_excel(all_rows, OUTPUT_FILE, verbose=False)

        browser.close()

    save_to_excel(all_rows, OUTPUT_FILE)
    print(f"\nГотово. Всего компаний без сайта: {len(all_rows)}")


if __name__ == "__main__":
    main()
