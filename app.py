# -*- coding: utf-8 -*-
"""Мини-CRM в духе amoCRM + парсер Яндекс.Карт.

Запуск:  python app.py   →  http://127.0.0.1:5000
При первом запуске сайт предложит создать администратора.

Облачная версия (Vercel) берёт настройки из переменных окружения:
DATABASE_URL (Postgres) и SECRET_KEY. Парсер в облаке не работает — его
запускают локально с тем же DATABASE_URL в файле .env, и найденные
компании сразу появляются в облачной CRM.
"""

import io
import os
import re
import secrets
import urllib.parse
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

BASE = Path(__file__).parent

# локальный .env (KEY=VALUE построчно) — до импорта db, который читает DATABASE_URL
_env = BASE / ".env"
if _env.exists():
    for _line in _env.read_text(encoding="utf-8").splitlines():
        if "=" in _line and not _line.lstrip().startswith("#"):
            _k, _v = _line.split("=", 1)
            os.environ.setdefault(_k.strip(), _v.strip().strip('"').strip("'"))

from flask import Flask, jsonify, redirect, render_template, request, send_file, session  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from werkzeug.security import check_password_hash, generate_password_hash  # noqa: E402

import db  # noqa: E402

IS_VERCEL = bool(os.environ.get("VERCEL"))


def _secret_key() -> str:
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    key_file = BASE / "secret.key"
    if not key_file.exists():
        key_file.write_text(secrets.token_hex(32))
    return key_file.read_text().strip()


app = Flask(__name__)
app.secret_key = _secret_key()
app.permanent_session_lifetime = timedelta(days=30)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_SECURE=IS_VERCEL)
# в облаке не трогаем задачи парсинга: их может прямо сейчас выполнять локальный парсер
db.init(stop_stale_jobs=not IS_VERCEL)

# парсеру нужен настоящий браузер (Edge) — он работает только в локальной версии
runner = None
if not IS_VERCEL:
    try:
        from scraper import Runner
        runner = Runner()
    except ImportError:  # playwright не установлен
        runner = None

CITY_PRESETS = [
    {"id": "10335", "slug": "tashkent", "name": "Ташкент"},
    {"id": "213", "slug": "moscow", "name": "Москва"},
    {"id": "2", "slug": "saint_petersburg", "name": "Санкт-Петербург"},
    {"id": "162", "slug": "almaty", "name": "Алматы"},
    {"id": "157", "slug": "minsk", "name": "Минск"},
]

LEAD_FIELDS = {
    "name": "Название", "budget": "Бюджет", "contact_person": "Контактное лицо", "phones": "Телефон",
    "email": "Email", "address": "Адрес", "city": "Город", "yandex_category": "Рубрика",
    "website": "Сайт", "tags": "Теги", "url": "Ссылка на Яндекс.Карты",
}


# ---------------- авторизация ----------------

def me():
    uid = session.get("uid")
    if not uid:
        return None
    u = db.one("SELECT id, name, login, role, active FROM users WHERE id=?", (uid,))
    return u if u and u["active"] else None


def login_required(fn):
    @wraps(fn)
    def wrap(*a, **kw):
        if not me():
            if request.path.startswith("/api/"):
                return jsonify(error="Нужно войти"), 401
            return redirect("/login")
        return fn(*a, **kw)
    return wrap


def admin_required(fn):
    @wraps(fn)
    @login_required
    def wrap(*a, **kw):
        if me()["role"] != "admin":
            return jsonify(error="Доступно только администратору"), 403
        return fn(*a, **kw)
    return wrap


def err(msg, code=400):
    return jsonify(error=msg), code


def body() -> dict:
    return request.get_json(force=True, silent=True) or {}


@app.route("/login", methods=["GET", "POST"])
def login():
    setup = db.one("SELECT COUNT(*) AS n FROM users")["n"] == 0
    if setup and IS_VERCEL:
        # иначе первый случайный посетитель публичного сайта стал бы администратором
        return render_template("login.html", setup=False, setup_blocked=True, error=None)
    error = None
    if request.method == "POST":
        f = request.form
        if setup:
            name, login_, pw = f.get("name", "").strip(), f.get("login", "").strip().lower(), f.get("password", "")
            if not name or not login_ or len(pw) < 6:
                error = "Заполните имя, логин и пароль (от 6 символов)"
            else:
                uid = db.insert("users", {"name": name, "login": login_, "pw_hash": generate_password_hash(pw),
                                          "role": "admin", "created_at": db.now()})
                # сделки из первой версии без ответственного — назначаем администратору
                db.run("UPDATE leads SET responsible_id=? WHERE responsible_id IS NULL", (uid,))
                session.permanent = True
                session["uid"] = uid
                return redirect("/")
        else:
            u = db.one("SELECT * FROM users WHERE login=?", (f.get("login", "").strip().lower(),))
            if u and u["active"] and check_password_hash(u["pw_hash"], f.get("password", "")):
                session.permanent = True
                session["uid"] = u["id"]
                return redirect("/")
            error = "Неверный логин или пароль"
    return render_template("login.html", setup=setup, error=error)


@app.get("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.get("/")
@login_required
def index():
    return render_template("app.html")


@app.get("/api/bootstrap")
@login_required
def bootstrap():
    return jsonify(
        me=me(),
        stages=db.stages(),
        users=db.all_("SELECT id, name, role, active FROM users ORDER BY name"),
        cities=[r["city"] for r in db.all_("SELECT DISTINCT city FROM leads WHERE city<>'' ORDER BY city")],
        task_kinds=db.TASK_KINDS,
        city_presets=CITY_PRESETS,
        parser_enabled=runner is not None,
    )


# ---------------- сделки ----------------

LEAD_SELECT = """
    SELECT l.*, s.name AS stage_name, s.color AS stage_color, s.kind AS stage_kind,
           u.name AS responsible_name,
           (SELECT MIN(due_at) FROM tasks t WHERE t.lead_id=l.id AND t.done=0) AS next_task
    FROM leads l
    LEFT JOIN stages s ON s.id=l.stage_id
    LEFT JOIN users u ON u.id=l.responsible_id
"""


def lead_filter(a) -> tuple[str, list]:
    where, args = [], []
    uid = session.get("uid")
    if a.get("stage_id", "").isdigit():
        where.append("l.stage_id=?"); args.append(int(a["stage_id"]))
    resp = a.get("responsible_id", "")
    if resp == "me":
        where.append("l.responsible_id=?"); args.append(uid)
    elif resp == "none":
        where.append("l.responsible_id IS NULL")
    elif resp.isdigit():
        where.append("l.responsible_id=?"); args.append(int(resp))
    if a.get("city"):
        where.append("l.city=?"); args.append(a["city"])
    if a.get("only_phone") == "1":
        where.append("COALESCE(l.phones,'')<>''")
    if a.get("job_id", "").isdigit():
        where.append("l.job_id=?"); args.append(int(a["job_id"]))
    task = a.get("task")
    nowts = db.now()[:16]
    if task == "none":
        where.append("NOT EXISTS (SELECT 1 FROM tasks t WHERE t.lead_id=l.id AND t.done=0)")
    elif task == "overdue":
        where.append("EXISTS (SELECT 1 FROM tasks t WHERE t.lead_id=l.id AND t.done=0 AND t.due_at<?)"); args.append(nowts)
    if a.get("active") == "1":
        where.append("COALESCE(s.kind,'open')='open'")
    q = (a.get("q") or "").strip()
    if q:
        where.append("(l.name LIKE ? OR l.phones LIKE ? OR l.address LIKE ? OR l.contact_person LIKE ? "
                     "OR l.tags LIKE ? OR l.yandex_category LIKE ? OR l.email LIKE ?)")
        args += [f"%{q}%"] * 7
    return (" WHERE " + " AND ".join(where)) if where else "", args


@app.get("/api/leads")
@login_required
def leads_list():
    where, args = lead_filter(request.args)
    total = db.one(f"SELECT COUNT(*) AS n FROM leads l LEFT JOIN stages s ON s.id=l.stage_id{where}", args)["n"]
    if request.args.get("view") == "kanban":
        items = db.all_(f"{LEAD_SELECT}{where} ORDER BY l.stage_changed_at DESC, l.id DESC LIMIT 3000", args)
        return jsonify(items=items, total=total)
    page = max(1, int(request.args.get("page", 1) or 1))
    size = 50
    sort = {"name": "l.name", "budget": "l.budget DESC", "created": "l.id DESC",
            "updated": "l.updated_at DESC"}.get(request.args.get("sort"), "l.id DESC")
    items = db.all_(f"{LEAD_SELECT}{where} ORDER BY {sort} LIMIT ? OFFSET ?", args + [size, (page - 1) * size])
    return jsonify(items=items, total=total, page=page, size=size)


@app.post("/api/leads")
@login_required
def lead_create():
    d = body()
    name = (d.get("name") or "").strip()
    if not name:
        return err("Укажите название сделки")
    stage_id = int(d.get("stage_id") or db.first_stage_id())
    ts = db.now()
    lead_id = db.insert("leads", {
        "name": name, "phones": (d.get("phones") or "").strip(), "budget": _int(d.get("budget")),
        "contact_person": (d.get("contact_person") or "").strip(), "stage_id": stage_id,
        "responsible_id": d.get("responsible_id") or session["uid"], "source": "Вручную",
        "city": "", "category_query": "", "yandex_category": "", "address": "", "rating": "",
        "reviews": "", "website": "", "url": "", "org_id": f"manual-{secrets.token_hex(6)}",
        "stage_changed_at": ts, "created_at": ts, "updated_at": ts,
    })
    db.add_event(lead_id, session["uid"], "system", "Сделка создана вручную")
    return jsonify(id=lead_id)


@app.get("/api/leads/<int:lead_id>")
@login_required
def lead_get(lead_id):
    lead = db.one(f"{LEAD_SELECT} WHERE l.id=?", (lead_id,))
    if not lead:
        return err("Сделка не найдена", 404)
    feed = db.all_(
        "SELECT e.*, u.name AS user_name, 'event' AS type FROM events e LEFT JOIN users u ON u.id=e.user_id "
        "WHERE e.lead_id=?", (lead_id,))
    tasks = db.all_(
        "SELECT t.*, u.name AS user_name, c.name AS author_name, 'task' AS type FROM tasks t "
        "LEFT JOIN users u ON u.id=t.user_id LEFT JOIN users c ON c.id=t.created_by WHERE t.lead_id=?", (lead_id,))
    return jsonify(lead=lead, events=feed, tasks=tasks)


def _int(v) -> int:
    try:
        return max(0, int(str(v).replace(" ", "").replace(" ", "") or 0))
    except ValueError:
        return 0


def apply_lead_changes(lead_id: int, d: dict, uid: int):
    lead = db.one("SELECT * FROM leads WHERE id=?", (lead_id,))
    if not lead:
        return False
    upd, ts = {}, db.now()
    for k, title in LEAD_FIELDS.items():
        if k in d:
            v = _int(d[k]) if k == "budget" else str(d[k] or "").strip()[:1000]
            if v != (lead[k] if lead[k] is not None else ("" if k != "budget" else 0)):
                upd[k] = v
                if k in ("budget", "name"):
                    db.add_event(lead_id, uid, "field", f"{title}: {lead[k] or '—'} → {v or '—'}")
    if "stage_id" in d and d["stage_id"] and int(d["stage_id"]) != lead["stage_id"]:
        new = db.one("SELECT * FROM stages WHERE id=?", (int(d["stage_id"]),))
        if new:
            old = db.one("SELECT name FROM stages WHERE id=?", (lead["stage_id"],))
            upd.update(stage_id=new["id"], stage_changed_at=ts,
                       closed_at=ts if new["kind"] in ("won", "lost") else None)
            db.add_event(lead_id, uid, "stage", f"Этап: {old['name'] if old else '—'} → {new['name']}")
    if "responsible_id" in d:
        rid = int(d["responsible_id"]) if d["responsible_id"] else None
        if rid != lead["responsible_id"]:
            u = db.one("SELECT name FROM users WHERE id=?", (rid,)) if rid else None
            upd["responsible_id"] = rid
            db.add_event(lead_id, uid, "field", f"Ответственный: {u['name'] if u else 'не назначен'}")
    if upd:
        upd["updated_at"] = ts
        db.update("leads", lead_id, upd)
    return True


@app.patch("/api/leads/<int:lead_id>")
@login_required
def lead_patch(lead_id):
    if not apply_lead_changes(lead_id, body(), session["uid"]):
        return err("Сделка не найдена", 404)
    return jsonify(lead=db.one(f"{LEAD_SELECT} WHERE l.id=?", (lead_id,)))


@app.delete("/api/leads/<int:lead_id>")
@login_required
def lead_delete(lead_id):
    db.run("DELETE FROM leads WHERE id=?", (lead_id,))
    return jsonify(ok=True)


@app.post("/api/leads/bulk")
@login_required
def leads_bulk():
    d = body()
    ids = [int(i) for i in d.get("ids", []) if str(i).isdigit()]
    action = d.get("action")
    for i in ids:
        if action == "delete":
            db.run("DELETE FROM leads WHERE id=?", (i,))
        elif action == "stage":
            apply_lead_changes(i, {"stage_id": d.get("value")}, session["uid"])
        elif action == "responsible":
            apply_lead_changes(i, {"responsible_id": d.get("value")}, session["uid"])
    return jsonify(ok=True, count=len(ids))


@app.post("/api/leads/<int:lead_id>/notes")
@login_required
def note_add(lead_id):
    text = (body().get("text") or "").strip()
    if not text:
        return err("Пустое примечание")
    db.add_event(lead_id, session["uid"], "note", text[:5000])
    db.run("UPDATE leads SET updated_at=? WHERE id=?", (db.now(), lead_id))
    return jsonify(ok=True)


@app.delete("/api/events/<int:event_id>")
@login_required
def event_delete(event_id):
    db.run("DELETE FROM events WHERE id=? AND kind='note' AND (user_id=? OR ?='admin')",
           (event_id, session["uid"], me()["role"]))
    return jsonify(ok=True)


# ---------------- задачи ----------------

def _due(v: str) -> str | None:
    v = (v or "").replace("T", " ")[:16]
    try:
        datetime.strptime(v, "%Y-%m-%d %H:%M")
        return v
    except ValueError:
        return None


@app.post("/api/leads/<int:lead_id>/tasks")
@login_required
def task_add(lead_id):
    d = body()
    due = _due(d.get("due_at"))
    if not due:
        return err("Укажите срок задачи")
    kind = d.get("kind") if d.get("kind") in db.TASK_KINDS else "call"
    db.insert("tasks", {
        "lead_id": lead_id, "user_id": int(d.get("user_id") or session["uid"]), "created_by": session["uid"],
        "kind": kind, "text": (d.get("text") or "").strip()[:2000], "due_at": due, "created_at": db.now(),
    })
    db.run("UPDATE leads SET updated_at=? WHERE id=?", (db.now(), lead_id))
    return jsonify(ok=True)


@app.get("/api/tasks")
@login_required
def tasks_list():
    a = request.args
    where, args = ["t.done=?"], [1 if a.get("done") == "1" else 0]
    user = a.get("user", "me")
    if user == "me":
        where.append("t.user_id=?"); args.append(session["uid"])
    elif user.isdigit():
        where.append("t.user_id=?"); args.append(int(user))
    order = "t.done_at DESC LIMIT 200" if a.get("done") == "1" else "t.due_at"
    return jsonify(db.all_(
        "SELECT t.*, l.name AS lead_name, l.phones AS lead_phones, u.name AS user_name, "
        "s.name AS stage_name, s.color AS stage_color FROM tasks t "
        "JOIN leads l ON l.id=t.lead_id LEFT JOIN users u ON u.id=t.user_id LEFT JOIN stages s ON s.id=l.stage_id "
        f"WHERE {' AND '.join(where)} ORDER BY {order}", args))


@app.patch("/api/tasks/<int:task_id>")
@login_required
def task_patch(task_id):
    d = body()
    t = db.one("SELECT * FROM tasks WHERE id=?", (task_id,))
    if not t:
        return err("Задача не найдена", 404)
    upd = {}
    if d.get("done") and not t["done"]:
        result = (d.get("result") or "").strip()[:2000]
        upd.update(done=1, result=result, done_at=db.now())
        text = f"Выполнена задача «{db.TASK_KINDS.get(t['kind'], 'Задача')}»" + (f": {t['text']}" if t["text"] else "")
        if result:
            text += f"\nРезультат: {result}"
        db.add_event(t["lead_id"], session["uid"], "task_done", text)
    if "due_at" in d:
        due = _due(d["due_at"])
        if not due:
            return err("Неверный срок")
        upd["due_at"] = due
    if "text" in d:
        upd["text"] = str(d["text"]).strip()[:2000]
    if "user_id" in d and str(d["user_id"]).isdigit():
        upd["user_id"] = int(d["user_id"])
    db.update("tasks", task_id, upd)
    return jsonify(ok=True)


@app.delete("/api/tasks/<int:task_id>")
@login_required
def task_delete(task_id):
    db.run("DELETE FROM tasks WHERE id=?", (task_id,))
    return jsonify(ok=True)


# ---------------- этапы воронки ----------------

@app.post("/api/stages")
@admin_required
def stage_add():
    d = body()
    name = (d.get("name") or "").strip()
    if not name:
        return err("Укажите название этапа")
    # новый открытый этап встаёт перед закрывающими (успех/отказ)
    pos = db.one("SELECT COALESCE(MAX(sort),0)+1 AS p FROM stages WHERE kind='open'")["p"]
    db.run("UPDATE stages SET sort=sort+1 WHERE sort>=?", (pos,))
    sid = db.insert("stages", {"name": name, "color": d.get("color") or "#8fd3ff", "sort": pos, "kind": "open"})
    return jsonify(id=sid)


@app.patch("/api/stages/<int:stage_id>")
@admin_required
def stage_patch(stage_id):
    d = body()
    upd = {k: str(d[k]).strip() for k in ("name", "color") if d.get(k)}
    db.update("stages", stage_id, upd)
    return jsonify(ok=True)


@app.put("/api/stages/order")
@admin_required
def stage_order():
    for i, sid in enumerate(body().get("ids", [])):
        db.run("UPDATE stages SET sort=? WHERE id=?", (i, int(sid)))
    return jsonify(ok=True)


@app.delete("/api/stages/<int:stage_id>")
@admin_required
def stage_delete(stage_id):
    s = db.one("SELECT * FROM stages WHERE id=?", (stage_id,))
    if not s:
        return err("Этап не найден", 404)
    if s["kind"] != "open":
        return err("Этапы «Успешно» и «Не реализовано» удалить нельзя")
    if db.one("SELECT COUNT(*) AS n FROM stages WHERE kind='open'")["n"] <= 1:
        return err("Должен остаться хотя бы один рабочий этап")
    db.run("DELETE FROM stages WHERE id=?", (stage_id,))
    db.run("UPDATE leads SET stage_id=? WHERE stage_id=?", (db.first_stage_id(), stage_id))
    return jsonify(ok=True)


# ---------------- пользователи ----------------

@app.get("/api/users")
@admin_required
def users_list():
    return jsonify(db.all_("SELECT id, name, login, role, active, created_at FROM users ORDER BY id"))


@app.post("/api/users")
@admin_required
def user_add():
    d = body()
    name, login_, pw = (d.get("name") or "").strip(), (d.get("login") or "").strip().lower(), d.get("password") or ""
    if not name or not login_ or len(pw) < 6:
        return err("Имя, логин и пароль (от 6 символов) обязательны")
    if db.one("SELECT 1 AS x FROM users WHERE login=?", (login_,)):
        return err("Такой логин уже есть")
    uid = db.insert("users", {"name": name, "login": login_, "pw_hash": generate_password_hash(pw),
                              "role": "admin" if d.get("role") == "admin" else "manager", "created_at": db.now()})
    return jsonify(id=uid)


@app.patch("/api/users/<int:user_id>")
@admin_required
def user_patch(user_id):
    d, upd = body(), {}
    if d.get("name"):
        upd["name"] = d["name"].strip()
    if d.get("password"):
        if len(d["password"]) < 6:
            return err("Пароль от 6 символов")
        upd["pw_hash"] = generate_password_hash(d["password"])
    if user_id == session["uid"] and (d.get("role") == "manager" or d.get("active") is False):
        return err("Нельзя понизить или отключить самого себя")
    if d.get("role") in ("admin", "manager"):
        upd["role"] = d["role"]
    if "active" in d:
        upd["active"] = 1 if d["active"] else 0
    db.update("users", user_id, upd)
    return jsonify(ok=True)


@app.post("/api/me/password")
@login_required
def my_password():
    d = body()
    u = db.one("SELECT * FROM users WHERE id=?", (session["uid"],))
    if not check_password_hash(u["pw_hash"], d.get("old", "")):
        return err("Текущий пароль неверный")
    if len(d.get("new", "")) < 6:
        return err("Новый пароль от 6 символов")
    db.update("users", u["id"], {"pw_hash": generate_password_hash(d["new"])})
    return jsonify(ok=True)


# ---------------- аналитика ----------------

@app.get("/api/analytics")
@login_required
def analytics():
    a = request.args
    today = db.today()
    d_from = a.get("from") or today.replace(day=1).isoformat()
    d_to = a.get("to") or today.isoformat()
    lo, hi = d_from + " 00:00:00", d_to + " 23:59:59"

    # CAST(... AS BIGINT): в Postgres SUM возвращает numeric, который ушёл бы в JSON строкой
    funnel = db.all_(
        "SELECT s.id, s.name, s.color, s.kind, COUNT(l.id) AS n, CAST(COALESCE(SUM(l.budget),0) AS BIGINT) AS budget "
        "FROM stages s LEFT JOIN leads l ON l.stage_id=s.id GROUP BY s.id ORDER BY s.sort, s.id")
    created = db.one("SELECT COUNT(*) AS n FROM leads WHERE created_at BETWEEN ? AND ?", (lo, hi))["n"]
    closed = db.all_(
        "SELECT s.kind, COUNT(*) AS n, CAST(COALESCE(SUM(l.budget),0) AS BIGINT) AS budget FROM leads l "
        "JOIN stages s ON s.id=l.stage_id WHERE s.kind IN ('won','lost') AND l.closed_at BETWEEN ? AND ? "
        "GROUP BY s.kind", (lo, hi))
    closed = {r["kind"]: r for r in closed}
    won, lost = closed.get("won", {"n": 0, "budget": 0}), closed.get("lost", {"n": 0, "budget": 0})

    managers = db.all_(
        """SELECT u.id, u.name,
             (SELECT COUNT(*) FROM leads l JOIN stages s ON s.id=l.stage_id
                WHERE l.responsible_id=u.id AND s.kind='open') AS in_work,
             (SELECT COUNT(*) FROM leads l JOIN stages s ON s.id=l.stage_id
                WHERE l.responsible_id=u.id AND s.kind='won' AND l.closed_at BETWEEN ? AND ?) AS won,
             (SELECT CAST(COALESCE(SUM(l.budget),0) AS BIGINT) FROM leads l JOIN stages s ON s.id=l.stage_id
                WHERE l.responsible_id=u.id AND s.kind='won' AND l.closed_at BETWEEN ? AND ?) AS won_budget,
             (SELECT COUNT(*) FROM leads l JOIN stages s ON s.id=l.stage_id
                WHERE l.responsible_id=u.id AND s.kind='lost' AND l.closed_at BETWEEN ? AND ?) AS lost,
             (SELECT COUNT(*) FROM tasks t WHERE t.user_id=u.id AND t.done=1 AND t.done_at BETWEEN ? AND ?) AS tasks_done,
             (SELECT COUNT(*) FROM tasks t WHERE t.user_id=u.id AND t.done=0 AND t.due_at < ?) AS tasks_overdue,
             (SELECT COUNT(*) FROM events e WHERE e.user_id=u.id AND e.kind='note' AND e.created_at BETWEEN ? AND ?) AS notes
           FROM users u WHERE u.active=1 ORDER BY won_budget DESC, won DESC""",
        (lo, hi) * 4 + (db.now()[:16], lo, hi))

    daily = db.all_(
        "SELECT substr(created_at,1,10) AS d, COUNT(*) AS n FROM leads WHERE created_at BETWEEN ? AND ? "
        "GROUP BY d ORDER BY d", (lo, hi))
    no_tasks = db.one(
        "SELECT COUNT(*) AS n FROM leads l JOIN stages s ON s.id=l.stage_id WHERE s.kind='open' "
        "AND NOT EXISTS (SELECT 1 FROM tasks t WHERE t.lead_id=l.id AND t.done=0)")["n"]
    return jsonify(
        period={"from": d_from, "to": d_to}, funnel=funnel, created=created, won=won, lost=lost,
        conversion=round(won["n"] * 100 / (won["n"] + lost["n"])) if won["n"] + lost["n"] else None,
        managers=managers, daily=daily, no_tasks=no_tasks,
    )


# ---------------- парсер ----------------

PARSER_OFF = "Парсер работает только в локальной версии (python app.py на компьютере с Edge)"


def _job_view(j: dict) -> dict:
    live = runner.get(j["id"]) if runner else None
    if live and j["status"] in ("queued", "running", "captcha"):
        j["status"] = live.status
    return j


@app.post("/api/jobs")
@login_required
def create_job():
    if not runner:
        return err(PARSER_OFF)
    d = body()
    categories = [c.strip() for c in re.split(r"[\n;]", d.get("categories", "")) if c.strip()]
    if not categories:
        return err("Укажите хотя бы одну категорию")
    link = (d.get("city_url") or "").strip()
    if link:
        m = re.search(r"/maps/(\d+)/([^/?#]+)", urllib.parse.unquote(link))
        if not m:
            return err("Не удалось распознать ссылку Яндекс.Карт (нужна вида yandex.ru/maps/<id>/<город>/...)")
        city = {"id": m.group(1), "slug": m.group(2), "name": (d.get("city_name") or m.group(2)).strip()}
    else:
        city = next((c for c in CITY_PRESETS if c["id"] == str(d.get("city_id"))), None)
    if not city:
        return err("Выберите город")
    try:
        max_results = max(1, min(int(d.get("max_results", 60)), 300))
    except (TypeError, ValueError):
        return err("Неверное число результатов")
    params = {"categories": categories, "city": city, "max_results": max_results,
              "only_no_site": bool(d.get("only_no_site", True))}
    return jsonify(id=runner.submit(params, user_id=session["uid"]))


@app.get("/api/jobs")
@login_required
def jobs():
    return jsonify([_job_view(j) for j in db.list_jobs()])


@app.get("/api/jobs/<int:job_id>")
@login_required
def job(job_id):
    j = db.get_job(job_id)
    if not j:
        return err("not found", 404)
    live = runner.get(job_id) if runner else None
    j["logs"] = list(live.logs) if live else []
    return jsonify(_job_view(j))


@app.post("/api/jobs/<int:job_id>/stop")
@login_required
def job_stop(job_id):
    if not runner:
        return err(PARSER_OFF)
    return jsonify(ok=runner.stop(job_id))


@app.post("/api/jobs/<int:job_id>/captcha")
@login_required
def job_captcha(job_id):
    if not runner:
        return err(PARSER_OFF)
    return jsonify(ok=runner.captcha_done(job_id))


# ---------------- экспорт ----------------

@app.get("/export.xlsx")
@login_required
def export():
    where, args = lead_filter(request.args)
    rows = db.all_(f"{LEAD_SELECT}{where} ORDER BY l.id DESC", args)
    wb = Workbook()
    ws = wb.active
    ws.title = "Сделки"
    headers = ["Название", "Этап", "Бюджет", "Ответственный", "Контактное лицо", "Телефон(ы)", "Email",
               "Город", "Адрес", "Рубрика (Яндекс)", "Рейтинг", "Отзывов", "Сайт", "Теги", "Источник",
               "Создана", "Ссылка на Яндекс.Карты"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["name"], r["stage_name"], r["budget"], r["responsible_name"], r["contact_person"],
                   r["phones"], r["email"], r["city"], r["address"], r["yandex_category"], r["rating"],
                   r["reviews"], r["website"], r["tags"], r["source"], r["created_at"], r["url"]])
    for i, w in enumerate([30, 20, 12, 18, 18, 24, 22, 14, 40, 24, 9, 9, 24, 16, 18, 18, 45], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"crm_{datetime.now():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    # use_reloader=False — иначе воркер парсера запустится дважды
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)
